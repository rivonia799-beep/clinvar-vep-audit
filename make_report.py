# -*- coding: utf-8 -*-
"""Build a one-page report as .docx and .pdf.

Run with a Python that has python-docx + reportlab + Pillow
(e.g. the Codex bundled runtime):
    python make_report.py

Outputs:
    REPORT_ONE_PAGER.docx      (editable; export to PDF from Word)
    REPORT_ONE_PAGER.pdf       (ready to send)
"""
import os

HERE = os.path.dirname(os.path.abspath(__file__))
FIG_A = os.path.join(HERE, "fig_A_auroc.png")
FIG_B = os.path.join(HERE, "fig_B_consequence_rule.png")

TITLE = "ClinVar VEP 评测集的分布污染诊断与指标归因分析"
ABSTRACT = (
    "在试跑 DNA 语言模型做变异效应预测（VEP）时，我怀疑通用 ClinVar 评测集存在分子后果分布污染，"
    "因此搭建了一条轻量诊断流水线，量化「按分子后果分层任务」与「扩大模型规模」各自对指标的贡献。"
)
FINDINGS = [
    ("1. 未分层集存在蛋白层面后果污染。",
     "在全量测试集上，36.0% 的正例为无义（终止获得）变异，而负例中为 0.0%。本数据集只含单碱基替换"
     "（SNV）——SNV 不会产生移码，但可通过提前引入终止密码子造成蛋白截短。这类变异的标签由蛋白层面"
     "后果决定；而模型输入是固定长度的裸基因组序列窗口，不含转录本结构与 CDS 阅读框信息，无法从输入"
     "推算终止后果。同时这类变异用后果注释规则（VEP / SnpEff）即可直接筛出，混入测试集后，基准不再"
     "反映调控序列建模能力。直接证据：一条不使用任何序列特征、仅依据 HGVS 后果标签的规则基线达到 "
     "AUPRC 0.4925，高于训练过的 4-mer 逻辑回归模型（0.2388），也高于 500M 参数的 Carbon-500M "
     "线性探针（0.4701）。"),
    ("2. 任务分层对指标的贡献大于模型规模。",
     "同一模型（4-mer 逻辑回归）下，将任务过滤到非蛋白质层变异：AUROC 0.663 → 0.807（+0.144）；"
     "同一任务下，将 4-mer 替换为 500M 参数 DNA 语言模型（Carbon-500M）：0.807 → 0.850（+0.043）。"
     "该增益在分层前后几乎一致（+0.048 vs +0.043），说明模型规模是近似可加的改进，无法补偿任务污染。"),
    ("3. AUPRC 的绝对值不可跨数据集比较。",
     "其随机基线等于正例比例。以「未分层集 → 非截断子集」这一步为例：正例比例由 0.111 降至 0.074，"
     "同一模型的 AUROC 基本不变（0.663 → 0.678），但 AUPRC 绝对值反而由 0.239 降至 0.231；"
     "而 AUPRC 相对随机基线的倍数却由 2.15× 升至 3.13×。可见绝对 AUPRC 会被基线下移拖累。"
     "评测应同时报告 AUROC。"),
]
METHOD = [
    ("数据", "ClinVar variant_summary，GRCh38，≥1 星评审，常染色体 SNV，P/LP vs B/LB（筛选出 1,423,685 条）"),
    ("切分", "按染色体硬切分：测试集 = 3, 6, 9, 12, 16, 18, 19, 21（与文献一致）"),
    ("样本", "训练 4,000 / 测试 2,000，按自然正负比分层；每条变异 ±100 bp（201 bp），来源 UCSC hg38"),
    ("评分器", "随机 / GC 含量 / 4-mer 上下文 / 4-mer 差分 / 后果规则基线 / Carbon-500M 线性探针"),
    ("模型", "逻辑回归（梯度下降，纯 NumPy 实现，无 scikit-learn 依赖）"),
]
FOOTER = "数据已冻结并附 MD5 校验和；python evaluate.py --all 可一键复现全部结果。本工作定位为评测诊断与归因分析（benchmark audit），非新基准。"


# ------------------------------------------------------------------ DOCX
def build_docx(path):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "微软雅黑"
    st.font.size = Pt(9)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.6)
        s.left_margin = s.right_margin = Cm(1.8)

    h = doc.add_paragraph()
    r = h.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(15)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.add_run("摘要　").bold = True
    p.add_run(ABSTRACT)
    p.paragraph_format.space_after = Pt(6)

    # two figures side by side
    t = doc.add_table(rows=1, cols=2)
    for cell, img in zip(t.rows[0].cells, (FIG_A, FIG_B)):
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.add_run().add_picture(img, width=Cm(8.2))
    doc.add_paragraph()

    for head, body in FINDINGS:
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(5)
        q.add_run(head + " ").bold = True
        q.add_run(body)

    doc.add_paragraph().add_run("方法要点").bold = True
    mt = doc.add_table(rows=len(METHOD), cols=2)
    mt.style = "Table Grid"
    for i, (k, v) in enumerate(METHOD):
        c0, c1 = mt.rows[i].cells
        c0.text = k
        c1.text = v
        c0.width = Cm(2.4)
        c1.width = Cm(14.4)
        for c in (c0, c1):
            for par in c.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(8.5)

    f = doc.add_paragraph()
    fr = f.add_run(FOOTER)
    fr.font.size = Pt(8)
    f.paragraph_format.space_before = Pt(6)

    doc.save(path)
    print("wrote", path)


# ------------------------------------------------------------------ PDF
def build_pdf(path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Image,
                                    Table, TableStyle)
    from reportlab.lib.styles import ParagraphStyle
    import os as _os

    # Embed real TTF fonts so the PDF renders anywhere (no viewer font pack needed)
    FDIR = r"C:\Windows\Fonts"
    REG, BLD = "Deng", "DengB"
    pdfmetrics.registerFont(TTFont(REG, _os.path.join(FDIR, "Deng.ttf")))
    pdfmetrics.registerFont(TTFont(BLD, _os.path.join(FDIR, "Dengb.ttf")))
    pdfmetrics.registerFontFamily(REG, normal=REG, bold=BLD, italic=REG, boldItalic=BLD)

    S_TITLE = ParagraphStyle("t", fontName=BLD, fontSize=16, leading=21,
                             alignment=1, spaceAfter=7)
    S_BODY = ParagraphStyle("b", fontName=REG, fontSize=10, leading=14.6, spaceAfter=5,
                        wordWrap="CJK")
    S_HEAD = ParagraphStyle("h", fontName=BLD, fontSize=12, leading=16,
                            spaceBefore=3, spaceAfter=3)
    S_SMALL = ParagraphStyle("s", fontName=REG, fontSize=9, leading=12.6, wordWrap="CJK")

    doc = SimpleDocTemplate(path, pagesize=A4,
                            leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=14 * mm, bottomMargin=12 * mm,
                            title=TITLE)
    W = A4[0] - 32 * mm
    fw = (W - 6 * mm) / 2
    fh = fw / 1.607

    flow = [Paragraph(TITLE, S_TITLE),
            Paragraph("<b>\u6458\u8981</b>\u3000" + ABSTRACT, S_BODY),
            Spacer(1, 4)]
    flow.append(Table([[Image(FIG_A, width=fw, height=fh),
                        Image(FIG_B, width=fw, height=fh)]],
                      colWidths=[fw + 3 * mm, fw + 3 * mm],
                      style=TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"),
                                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE")])))
    flow.append(Spacer(1, 8))
    for head, body in FINDINGS:
        flow.append(Paragraph(f"<b>{head}</b> {body}", S_BODY))
    flow.append(Spacer(1, 4))
    flow.append(Paragraph("<b>\u65b9\u6cd5\u8981\u70b9</b>", S_HEAD))
    data = [[Paragraph(f"<b>{k}</b>", S_SMALL), Paragraph(v, S_SMALL)] for k, v in METHOD]
    t = Table(data, colWidths=[20 * mm, W - 20 * mm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#bbbbbb")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("LEFTPADDING", (0, 0), (-1, -1), 4),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                           ("TOPPADDING", (0, 0), (-1, -1), 3),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    flow.append(t)
    flow.append(Spacer(1, 8))
    flow.append(Paragraph(FOOTER, ParagraphStyle(
        "f", fontName=REG, fontSize=8.5, leading=12,
        textColor=colors.HexColor("#555555"))))
    doc.build(flow)
    print("wrote", path)


if __name__ == "__main__":
    build_docx(os.path.join(HERE, "REPORT_ONE_PAGER.docx"))
    build_pdf(os.path.join(HERE, "REPORT_ONE_PAGER.pdf"))
